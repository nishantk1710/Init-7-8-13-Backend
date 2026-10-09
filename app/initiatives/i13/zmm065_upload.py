"""Monthly ZMM065 uploads: parse, store, and pick the reference validation reads.

FR-6 asks for a monthly reconciliation against ZMM065. The report reaches the
platform as a workbook export, and from August 2026 VZI uploads it on the
Validation screen each month rather than through the seed loader.

Parsing
-------
The two delivered workbooks are shaped differently -- BMM's data is ``Sheet1``
of five, GB's is ``Sheet2`` of three, both with a title row above the header --
and nothing guarantees next month's sheet is named the same. So the data sheet
is found by its *columns*, not its name: the first sheet whose header (within
the first few rows) carries every column validation reads. Headers are
sanitised exactly as the seed loader does (``app.seed.reader.column_name``), so
an upload and the seeded July table agree on what a column is called.

One upload is one plant -- ZMM065 is run per site. The plant is read from the
rows, not chosen on a form, so a BMM file cannot be filed as Gamsberg. Rows for
a plant outside the platform's scope are skipped and counted.

Which reference validation reads
--------------------------------
Without a month: per plant, the latest upload (by report month, then upload
time); a plant with no upload falls back to its seeded ``raw_zmm065_*`` table,
so today's July data keeps working. With a month: only that month's uploads,
and no fallback -- reconciling August's numbers against July's report would be
answering a different question.

Uploads are appended, never overwritten. Uploading the same plant and month
again requires ``replace=True`` and adds a newer row; the earlier upload stays
as the record of what was reconciled before.
"""

from __future__ import annotations

from collections import Counter
from collections.abc import Iterable
from dataclasses import dataclass
from datetime import date, datetime
from typing import IO, Literal

import openpyxl
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.initiatives.i13.report_validation import zmm065_report_date
from app.integrations.sap.postgres_reports import Zmm065Row, fetch_zmm065_rows, key, parse_date, parse_int
from app.models.i13_zmm065_upload import Zmm065Upload, Zmm065UploadRow
from app.seed.reader import to_text, unique_column_names
from app.shared.plant_scope import IN_SCOPE_PLANTS, normalise

#: The columns validation reads. A sheet without all of them is not the data.
REQUIRED_COLUMNS: tuple[str, ...] = ("mat_code", "plant", "stock_type", "last_gi_dt", "days")

#: How far down a sheet to look for the header row. Both delivered workbooks
#: carry one title row above it.
HEADER_SEARCH_ROWS = 10

#: Rows are inserted in batches of this size.
INSERT_BATCH = 2000


class Zmm065UploadError(ValueError):
    """The workbook cannot be accepted. The message is shown to the uploader."""


class Zmm065DuplicateUpload(Zmm065UploadError):
    """A report for this plant and month is already uploaded."""

    def __init__(self, existing: Zmm065Upload) -> None:
        self.existing = existing
        super().__init__(
            f"A ZMM065 report for plant {existing.plant}, {existing.report_month:%B %Y}, was already "
            f"uploaded on {existing.uploaded_at:%d %b %Y} by {existing.uploaded_by} "
            f"({existing.file_name}). Upload again with 'replace' to make this file the current one; "
            "the earlier upload is kept."
        )


@dataclass(frozen=True)
class ParsedZmm065:
    sheet_name: str
    plant: str
    rows: tuple[Zmm065Row, ...]
    #: Rows skipped because their plant is outside the platform's scope.
    skipped_out_of_scope: int
    #: Rows skipped because material or plant was blank (footers, subtotals).
    skipped_blank: int


@dataclass(frozen=True)
class Zmm065Source:
    """Where one plant's ZMM065 rows came from, for the Validation screen."""

    plant: str
    source: Literal["UPLOAD", "SEED"]
    row_count: int
    upload_id: int | None = None
    report_month: date | None = None
    report_date: date | None = None
    file_name: str | None = None
    uploaded_by: str | None = None
    uploaded_at: datetime | None = None


def parse_month(raw: str) -> date:
    """``2026-08`` (or ``2026-08-01``) -> the first of that month."""
    text = (raw or "").strip()
    try:
        year, month = (int(part) for part in text.split("-")[:2])
        return date(year, month, 1)
    except ValueError:
        raise Zmm065UploadError(f"Report month {raw!r} is not a month; expected YYYY-MM.") from None


def parse_zmm065_workbook(handle: IO[bytes]) -> ParsedZmm065:
    """Read one ZMM065 workbook. Raises :class:`Zmm065UploadError` with a
    message fit for the uploader when it cannot be used."""
    try:
        workbook = openpyxl.load_workbook(handle, read_only=True, data_only=True)
    except Exception as error:  # openpyxl raises several unrelated types for a non-workbook
        raise Zmm065UploadError(f"The file could not be read as an Excel workbook (.xlsx): {error}") from None
    try:
        for worksheet in workbook.worksheets:
            header = _find_header(worksheet)
            if header is not None:
                header_row, names = header
                return _read_rows(worksheet, header_row, names)
        raise Zmm065UploadError(
            "No sheet in this workbook has the ZMM065 columns validation needs "
            f"({', '.join(REQUIRED_COLUMNS)}). Sheets present: {', '.join(workbook.sheetnames)}."
        )
    finally:
        workbook.close()


def _find_header(worksheet) -> tuple[int, list[str]] | None:
    for number, row in enumerate(
        worksheet.iter_rows(min_row=1, max_row=HEADER_SEARCH_ROWS, values_only=True), start=1
    ):
        names = unique_column_names(list(row))
        if all(column in names for column in REQUIRED_COLUMNS):
            return number, names
    return None


def _read_rows(worksheet, header_row: int, names: list[str]) -> ParsedZmm065:
    position = {column: names.index(column) for column in REQUIRED_COLUMNS}
    rows: list[Zmm065Row] = []
    out_of_scope = 0
    blank = 0
    for raw in worksheet.iter_rows(min_row=header_row + 1, values_only=True):
        values = {column: to_text(raw[i]) if i < len(raw) else None for column, i in position.items()}
        material = key(values["mat_code"])
        plant = normalise(values["plant"])
        if not material or not plant:
            blank += 1
            continue
        if plant not in IN_SCOPE_PLANTS:
            out_of_scope += 1
            continue
        rows.append(
            Zmm065Row(
                material=material,
                plant=plant,
                stock_type=(values["stock_type"] or "").strip(),
                last_gi_date=parse_date(values["last_gi_dt"]),
                days=parse_int(values["days"]),
            )
        )

    plants = Counter(row.plant for row in rows)
    if not plants:
        raise Zmm065UploadError(
            f"Sheet {worksheet.title!r} has no rows for the in-scope plants "
            f"({', '.join(IN_SCOPE_PLANTS)}); {out_of_scope} row(s) were for other plants."
        )
    if len(plants) > 1:
        found = ", ".join(f"{plant} ({count} rows)" for plant, count in sorted(plants.items()))
        raise Zmm065UploadError(
            f"ZMM065 is uploaded one plant at a time, and this file holds rows for {found}. "
            "Upload each site's report separately."
        )
    return ParsedZmm065(
        sheet_name=worksheet.title,
        plant=next(iter(plants)),
        rows=tuple(rows),
        skipped_out_of_scope=out_of_scope,
        skipped_blank=blank,
    )


def store_upload(
    db: Session,
    parsed: ParsedZmm065,
    *,
    report_month: date,
    file_name: str,
    uploaded_by: str,
    replace: bool = False,
) -> tuple[Zmm065Upload, bool]:
    """Persist one parsed upload and commit. Refuses a second upload for the
    same plant and month unless ``replace`` is set. Returns the upload and
    whether it superseded an earlier one."""
    existing = db.scalars(
        select(Zmm065Upload)
        .where(Zmm065Upload.plant == parsed.plant, Zmm065Upload.report_month == report_month)
        .order_by(Zmm065Upload.uploaded_at.desc(), Zmm065Upload.id.desc())
    ).first()
    if existing is not None and not replace:
        raise Zmm065DuplicateUpload(existing)

    upload = Zmm065Upload(
        plant=parsed.plant,
        report_month=report_month,
        report_date=zmm065_report_date(list(parsed.rows)),
        file_name=file_name[:255],
        sheet_name=parsed.sheet_name[:64],
        row_count=len(parsed.rows),
        uploaded_by=uploaded_by[:128],
    )
    db.add(upload)
    db.flush()
    for start in range(0, len(parsed.rows), INSERT_BATCH):
        db.add_all(
            Zmm065UploadRow(
                upload_id=upload.id,
                material=row.material[:40],
                plant=row.plant,
                stock_type=row.stock_type[:40],
                last_gi_date=row.last_gi_date,
                days=row.days,
            )
            for row in parsed.rows[start : start + INSERT_BATCH]
        )
        db.flush()
    db.commit()
    db.refresh(upload)
    return upload, existing is not None


def list_uploads(db: Session) -> list[Zmm065Upload]:
    """Every upload, newest month first, then newest upload first."""
    return list(
        db.scalars(
            select(Zmm065Upload).order_by(
                Zmm065Upload.report_month.desc(), Zmm065Upload.uploaded_at.desc(), Zmm065Upload.id.desc()
            )
        )
    )


def current_uploads(uploads: Iterable[Zmm065Upload], report_month: date | None = None) -> dict[str, Zmm065Upload]:
    """Per plant, the upload validation reads. ``uploads`` must be in
    :func:`list_uploads` order, so the first one seen for a plant wins."""
    chosen: dict[str, Zmm065Upload] = {}
    for upload in uploads:
        if report_month is not None and upload.report_month != report_month:
            continue
        chosen.setdefault(upload.plant, upload)
    return chosen


def zmm065_reference(
    db: Session, report_month: date | None = None
) -> tuple[list[Zmm065Row] | None, list[Zmm065Source]]:
    """The ZMM065 rows validation reconciles against, and where each plant's
    came from. Rows are ``None`` when nothing at all is available -- the same
    "reference missing" answer ``fetch_zmm065_rows`` gives."""
    chosen = current_uploads(list_uploads(db), report_month)
    rows: list[Zmm065Row] = []
    sources: list[Zmm065Source] = []
    for plant in sorted(chosen):
        upload = chosen[plant]
        records = db.scalars(select(Zmm065UploadRow).where(Zmm065UploadRow.upload_id == upload.id))
        rows.extend(
            Zmm065Row(
                material=r.material,
                plant=r.plant,
                stock_type=r.stock_type,
                last_gi_date=r.last_gi_date,
                days=r.days,
            )
            for r in records
        )
        sources.append(
            Zmm065Source(
                plant=plant,
                source="UPLOAD",
                row_count=upload.row_count,
                upload_id=upload.id,
                report_month=upload.report_month,
                report_date=upload.report_date,
                file_name=upload.file_name,
                uploaded_by=upload.uploaded_by,
                uploaded_at=upload.uploaded_at,
            )
        )

    if report_month is None:
        seeded = fetch_zmm065_rows(db)
        if seeded is not None:
            fallback = [row for row in seeded if row.plant not in chosen]
            rows.extend(fallback)
            for plant, count in sorted(Counter(row.plant for row in fallback).items()):
                plant_rows = [row for row in fallback if row.plant == plant]
                sources.append(
                    Zmm065Source(
                        plant=plant,
                        source="SEED",
                        row_count=count,
                        report_date=zmm065_report_date(plant_rows),
                    )
                )

    return (rows if sources else None), sources
