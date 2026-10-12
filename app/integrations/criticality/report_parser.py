"""Read one plant's monthly criticality report (ZMM065).

Pure: rows in, rows and counts out. Nothing here touches the database, so every
rule below is testable on a workbook built in memory.

Finding the data
----------------
The two July workbooks are shaped differently -- BMM's data is ``Sheet1`` of
five (the others are pivots), GB's is ``Sheet2`` of three, both with an empty
row above the header, and GB carries an extra ``Sl.No`` column first. Nothing
guarantees next month's sheet keeps its name or position. So the data sheet is
found by its *columns*: the first sheet whose header, within the first
``HEADER_SEARCH_ROWS`` rows, carries ``mat_code``, ``plant`` and
``criticality``. Headers are sanitised exactly as the seed loader does
(``app.seed.reader.unique_column_names``), so ``Mat.Code`` is ``mat_code`` here
and in ``raw_zmm065_*`` alike.

What an upload refuses, and why
-------------------------------
* **Another plant's rows.** The uploader picks the plant; the file must agree.
  Filing Gamsberg's report as Black Mountain would silently re-tier ten
  thousand materials at the wrong site.
* **A tier that is not one of the five.** ``parse_tier`` returns ``None`` for
  it, and storing that as "blank" would quietly drop a material's criticality.
  Criticality drives the service level; a typo has to be seen.
* **The same material twice with two different tiers.** There is no correct
  winner inside one report.

What is accepted and counted: blank tiers (stored as no tier -- "not maintained
this month", which does not erase an earlier month's tier), the same material
repeated with the same tier (collapsed), and rows without a material number or
plant (footers, subtotals).

The July delivery, already loaded into ``raw_zmm065_*``, goes through the same
rules in *lenient* mode (``strict=False``): it was delivered, not chosen, so
nothing in it can be refused -- other plants' rows are dropped, an unknown tier
is read as blank and the first of two conflicting rows is kept, all counted in
``ParsedReport.issues`` so the import can log them.
"""

from __future__ import annotations

import warnings
from collections import Counter
from dataclasses import dataclass, field
from datetime import date
from typing import IO

import openpyxl

from app.core.criticality import parse_tier
from app.seed.reader import to_text, unique_column_names
from app.shared.plant_scope import normalise as normalise_plant

#: Columns a sheet must have to be the report.
REQUIRED_COLUMNS: tuple[str, ...] = ("mat_code", "plant", "criticality")

#: Optional columns, by every sanitised header they have been seen under.
#: ``material_description`` is the July name; the others are the truncations
#: an export is known to produce.
OPTIONAL_COLUMNS: dict[str, tuple[str, ...]] = {
    "description": ("material_description", "material_descrip", "material_desc", "description"),
    "mattyp": ("mattyp", "mat_typ", "material_type"),
    "mrp_type": ("mrp_type", "mrptype"),
}

#: How far down a sheet to look for the header row.
HEADER_SEARCH_ROWS = 10

#: How many offending values an error message names before summarising.
_EXAMPLES = 10

_LIMITS = {"material": 40, "description": 200, "mattyp": 10, "mrp_type": 10}


class CriticalityReportError(ValueError):
    """The report cannot be accepted. The message is shown to the uploader."""


@dataclass(frozen=True)
class ReportRow:
    material: str
    material_raw: str
    plant: str
    criticality: str | None
    description: str | None = None
    mattyp: str | None = None
    mrp_type: str | None = None


@dataclass(frozen=True)
class ParsedReport:
    sheet_name: str
    plant: str
    #: One per material, in file order.
    rows: tuple[ReportRow, ...]
    #: Rows that carried a material number and plant, before duplicates were
    #: collapsed.
    rows_read: int
    #: Repeats of a material already read, collapsed into the first.
    duplicate_rows: int
    #: Materials whose tier is blank.
    blank_tiers: int
    #: Rows with no material number or no plant -- footers, subtotals.
    skipped_blank_material: int
    tiers: Counter = field(default_factory=Counter)
    #: Lenient mode only: what strict mode would have refused.
    issues: tuple[str, ...] = ()


def material_key(raw: str | None) -> str:
    """Leading zeros stripped, trimmed -- the form the ``n_*`` views and
    ``sap_key`` use. All zeros stays as typed rather than becoming empty."""
    value = (raw or "").strip()
    return value.lstrip("0") or value


def parse_month(raw: str) -> date:
    """``2026-08`` (or ``2026-08-15``) -> the first of that month."""
    text = (raw or "").strip()
    try:
        year, month = (int(part) for part in text.split("-")[:2])
        return date(year, month, 1)
    except ValueError:
        raise CriticalityReportError(f"Report month {raw!r} is not a month; expected YYYY-MM.") from None


def parse_report(handle: IO[bytes], *, plant: str) -> ParsedReport:
    """Read one uploaded report for ``plant``, strictly. Raises
    :class:`CriticalityReportError` with a message fit for the uploader."""
    try:
        workbook = openpyxl.load_workbook(handle, read_only=True, data_only=True)
    except Exception as error:  # openpyxl raises several unrelated types for a non-workbook
        raise CriticalityReportError(f"The file could not be read as an Excel workbook (.xlsx): {error}") from None
    # The July BMM export carries date-formatted cells holding serial -693595
    # in columns this never reads; openpyxl warns once per cell, which would
    # bury the server log on every upload.
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")
        return _parse_workbook(workbook, plant)


def _parse_workbook(workbook, plant: str) -> ParsedReport:
    try:
        for worksheet in workbook.worksheets:
            header = _find_header(worksheet)
            if header is None:
                continue
            header_row, names = header
            position = {column: names.index(column) for column in REQUIRED_COLUMNS}
            for column, aliases in OPTIONAL_COLUMNS.items():
                found = next((alias for alias in aliases if alias in names), None)
                if found is not None:
                    position[column] = names.index(found)
            collector = RowCollector(plant, strict=True)
            for raw in worksheet.iter_rows(min_row=header_row + 1, values_only=True):
                values = {column: to_text(raw[i]) if i < len(raw) else None for column, i in position.items()}
                collector.add(
                    values["mat_code"],
                    values["plant"],
                    values["criticality"],
                    description=values.get("description"),
                    mattyp=values.get("mattyp"),
                    mrp_type=values.get("mrp_type"),
                )
            return collector.finish(worksheet.title)
        raise CriticalityReportError(
            "No sheet in this workbook has the columns Mat.Code, Plant and Criticality. "
            f"Sheets present: {', '.join(workbook.sheetnames)}."
        )
    finally:
        workbook.close()


def _find_header(worksheet) -> tuple[int, list[str]] | None:
    for number, row in enumerate(
        worksheet.iter_rows(min_row=1, max_row=HEADER_SEARCH_ROWS, values_only=True), start=1
    ):
        names = unique_column_names(list(row))
        if all(column in names for column in REQUIRED_COLUMNS):
            return number, names
    return None


def _clip(value: str | None, name: str) -> str | None:
    if value is None:
        return None
    value = value.strip()
    return value[: _LIMITS[name]] or None


class RowCollector:
    """Applies the row rules one row at a time, so a workbook and an
    already-loaded table are read by the same code."""

    def __init__(self, plant: str, *, strict: bool) -> None:
        self.plant = plant
        self.strict = strict
        self._by_material: dict[str, ReportRow] = {}
        self._other_plants: Counter = Counter()
        self._unknown_tiers: Counter = Counter()
        self._conflicts: dict[str, set[str]] = {}
        self._rows_read = 0
        self._duplicates = 0
        self._skipped = 0

    def add(
        self,
        material_raw: str | None,
        plant: str | None,
        tier_text: str | None,
        *,
        description: str | None = None,
        mattyp: str | None = None,
        mrp_type: str | None = None,
    ) -> None:
        material_raw = (material_raw or "").strip()
        row_plant = normalise_plant(plant)
        if not material_raw or not row_plant:
            self._skipped += 1
            return
        self._rows_read += 1
        if row_plant != self.plant:
            self._other_plants[row_plant] += 1
            return

        tier = parse_tier(tier_text)
        if tier is None and (tier_text or "").strip():
            self._unknown_tiers[tier_text.strip()] += 1
            if self.strict:
                return

        row = ReportRow(
            material=material_key(material_raw)[: _LIMITS["material"]],
            material_raw=material_raw[: _LIMITS["material"]],
            plant=row_plant,
            criticality=tier.value if tier is not None else None,
            description=_clip(description, "description"),
            mattyp=_clip(mattyp, "mattyp"),
            mrp_type=_clip(mrp_type, "mrp_type"),
        )
        earlier = self._by_material.get(row.material)
        if earlier is None:
            self._by_material[row.material] = row
            return
        self._duplicates += 1
        if earlier.criticality and row.criticality and earlier.criticality != row.criticality:
            self._conflicts.setdefault(row.material, {earlier.criticality}).add(row.criticality)
        elif earlier.criticality is None and row.criticality is not None:
            # A blank first occurrence and a tiered repeat: the repeat says more.
            self._by_material[row.material] = row

    def finish(self, sheet_name: str) -> ParsedReport:
        problems = self._problems()
        if self.strict and problems:
            raise CriticalityReportError(problems[0])
        if self.strict and not self._by_material:
            raise CriticalityReportError(f"Sheet {sheet_name!r} has no material rows for plant {self.plant}.")

        rows = tuple(self._by_material.values())
        return ParsedReport(
            sheet_name=sheet_name,
            plant=self.plant,
            rows=rows,
            rows_read=self._rows_read,
            duplicate_rows=self._duplicates,
            blank_tiers=sum(1 for row in rows if row.criticality is None),
            skipped_blank_material=self._skipped,
            tiers=Counter(row.criticality for row in rows if row.criticality),
            issues=() if self.strict else tuple(problems),
        )

    def _problems(self) -> list[str]:
        problems: list[str] = []
        if self._other_plants:
            found = ", ".join(f"{code} ({count:,} rows)" for code, count in sorted(self._other_plants.items()))
            if not self._by_material and not self._unknown_tiers:
                problems.append(
                    f"This file is for plant {found}, but you picked {self.plant}. "
                    "Pick the plant the report was run for."
                )
            else:
                problems.append(
                    f"This file holds rows for plant {found} as well as {self.plant}. "
                    "Each plant's report is uploaded on its own."
                )
        if self._unknown_tiers:
            listed = ", ".join(
                f"{value!r} ({count})" for value, count in self._unknown_tiers.most_common(_EXAMPLES)
            )
            problems.append(
                "Criticality must be NORMAL, OBSOLETE, CRITICAL, IMPACT or INSURANCE (or blank). "
                f"Unrecognised: {listed}."
            )
        if self._conflicts:
            listed = "; ".join(
                f"{material} ({' / '.join(sorted(tiers))})"
                for material, tiers in list(self._conflicts.items())[:_EXAMPLES]
            )
            more = f" and {len(self._conflicts) - _EXAMPLES} more" if len(self._conflicts) > _EXAMPLES else ""
            problems.append(
                f"{len(self._conflicts)} material(s) appear more than once with different tiers: {listed}{more}."
            )
        return problems
