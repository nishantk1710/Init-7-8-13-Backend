"""The I13 summary report: facts from the snapshot, words that may not invent.

What these guard against is mostly silent: a model-written figure looks like a
figure. So the number rule is tested from both sides, the fallback is tested
for every way a model can fail, and the routes are tested switched off as well
as on -- off is how the feature ships.

No database server and no network: SQLite in memory for the two tables the
report touches, a stand-in snapshot, and the provider monkeypatched at the same
seam ``test_quantity_suggestion_reason.py`` uses.
"""

from __future__ import annotations

import json
from datetime import date, datetime, timezone
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.api.i13.deps import require_snapshot
from app.core.ai import AITimeoutError, Completion, TokenUsage
from app.core.config import Settings, get_settings
from app.core.db import get_db
from app.initiatives.i13.models import AcquiredVsPlanStatus
from app.initiatives.i13.report import facts as facts_module
from app.initiatives.i13.report.export import build_report_workbook
from app.initiatives.i13.report.facts import build_report_facts
from app.initiatives.i13.report.labels import plant_label
from app.initiatives.i13.report.narrative import (
    DETERMINISTIC,
    MODEL,
    PROMPT_ID,
    invented_numbers,
    leaked_codes,
    parse_reply,
    template_narrative,
    write_narrative,
)
from app.initiatives.i13.summary import I13Summary
from app.main import create_app
from app.models.base import Base
from app.models.i13_act_exception import ActExceptionRecord
from app.models.i13_summary_report import I13SummaryReport
from app.shared.material_scope import MaterialScope

TODAY = date(2026, 10, 13)
OAR = MaterialScope.OAR


# --- a stand-in snapshot and database ----------------------------------------------


def _watch(material: str, plant: str, status: AcquiredVsPlanStatus, scope=OAR):
    return SimpleNamespace(material=material, plant=plant, material_scope=scope, acquired_vs_plan_status=status)


def _grni(material: str, plant: str, days: int, outstanding: str = "4"):
    ledger = SimpleNamespace(
        material=material, plant=plant, material_scope=OAR,
        reservation_number="0012345", reservation_item="0001", po_number="4500000001",
    )
    return SimpleNamespace(ledger=ledger, outstanding_quantity=Decimal(outstanding), days_since_gr=days)


SNAPSHOT = SimpleNamespace(
    version=12,
    reference_date=TODAY,
    built_at=datetime(2026, 10, 13, 2, 0, tzinfo=timezone.utc),
    watch={
        ("M1", "1300"): _watch("M1", "1300", AcquiredVsPlanStatus.NO_PLAN),
        ("M2", "1300"): _watch("M2", "1300", AcquiredVsPlanStatus.ABOVE_PLAN),
        ("M3", "1500"): _watch("M3", "1500", AcquiredVsPlanStatus.NO_PLAN),
        ("M4", "1500"): _watch("M4", "1500", AcquiredVsPlanStatus.ON_PLAN, scope=MaterialScope.MIN_MAX),
    },
    grni_entries=[_grni("M1", "1300", 140), _grni("M3", "1500", 61)],
)

SUMMARY = I13Summary(
    total_oar_positions=1200,
    fast_moving_count=300,
    slow_moving_count=150,
    non_moving_count=750,
    gr_not_issued_30_day_count=2,
    plan_breach_count=18,
    no_plan_count=640,
    reclassification_candidate_count=1,
    valuation_is_mocked=True,
    reference_plan_count=0,
)


@pytest.fixture
def factory():
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine, tables=[ActExceptionRecord.__table__, I13SummaryReport.__table__])
    return sessionmaker(bind=engine)


def _exception(session, exception_id: str, status: str, kind: str, plant: str, detected: datetime) -> None:
    session.add(
        ActExceptionRecord(
            exception_id=exception_id, exception_type=kind, status=status, material="M1", plant=plant,
            reservation_number="0012345", reservation_item="0001", detected_at=detected,
            reason="test", evidence_json="{}",
        )
    )


@pytest.fixture
def db(factory):
    with factory() as session:
        _exception(session, "E1", "ESCALATED", "NO_PLAN", "1300", datetime(2026, 8, 1, tzinfo=timezone.utc))
        _exception(session, "E2", "OPEN", "PLAN_BREACH", "1300", datetime(2026, 9, 1, tzinfo=timezone.utc))
        _exception(session, "E3", "RESOLVED", "NO_PLAN", "1500", datetime(2026, 7, 1, tzinfo=timezone.utc))
        session.commit()
        yield session


@pytest.fixture(autouse=True)
def _no_side_reads(monkeypatch: pytest.MonkeyPatch):
    """The summary and the material-name list both read the real database;
    neither belongs in a unit test."""
    monkeypatch.setattr(facts_module, "_summary", lambda snapshot, db, plant: SUMMARY)
    monkeypatch.setattr(facts_module, "_names", lambda materials: {})


@pytest.fixture
def facts(db):
    return build_report_facts(SNAPSHOT, db, today=TODAY)


# --- facts -----------------------------------------------------------------------------


def test_facts_carry_the_summary_and_its_shares(facts) -> None:
    assert facts["positions"]["total"] == 1200
    assert facts["positions"]["NON_MOVING"] == 750
    assert facts["positions"]["non_moving_pct"] == 62.5
    assert facts["valuation_is_mocked"] is True
    assert facts["data_as_of"] == "2026-10-13"


def test_acquired_vs_plan_counts_oar_positions_only(facts) -> None:
    counts = facts["plans"]["acquired_vs_plan"]
    assert counts["NO_PLAN"] == 2
    assert counts["ABOVE_PLAN"] == 1
    assert counts["ON_PLAN"] == 0, "the Min-Max position is not an OAR position"


def test_exceptions_are_grouped_and_the_oldest_unresolved_listed(facts) -> None:
    exc = facts["exceptions"]
    assert exc["total"] == 3
    assert exc["unresolved"] == 2
    assert exc["by_status"]["ESCALATED"] == 1
    assert exc["by_type"]["NO_PLAN"] == 2
    assert [e["exception_id"] for e in exc["oldest_unresolved"]] == ["E1", "E2"], "resolved E3 is not listed"
    assert exc["oldest_unresolved"][0]["days_open"] == 73


def test_one_plant_scopes_every_section(db) -> None:
    facts = build_report_facts(SNAPSHOT, db, plant="1500", today=TODAY)
    assert facts["scope"]["label"] == plant_label("1500")
    assert "(1500)" in facts["scope"]["label"], "the screens show a plant as name and code"
    assert facts["plans"]["acquired_vs_plan"]["NO_PLAN"] == 1
    assert facts["gr_not_issued"]["open_line_count"] == 1
    assert facts["exceptions"]["total"] == 1
    assert facts["gr_not_issued"]["longest_waiting"][0]["plant"] == plant_label("1500")


def test_facts_are_json_ready(facts) -> None:
    json.dumps(facts)


# --- the number rule ---------------------------------------------------------------


def _reply(**overrides) -> str:
    body = {
        "headline": "750 of 1,200 OAR positions (62.5%) are non-moving.",
        "findings": ["18 plan breaches.", "Reservation 0012345/0001 for M1 has waited 140 days since GR."],
        "risks": ["Idle stock."],
        "actions": [{"action": "Chase the 2 oldest received lines.", "owner": "Plant stores"}],
    }
    body.update(overrides)
    return json.dumps(body)


def test_numbers_from_the_facts_pass(facts) -> None:
    assert invented_numbers(parse_reply(_reply()), facts) == set()


def test_a_number_the_facts_do_not_hold_is_caught(facts) -> None:
    narrative = parse_reply(_reply(headline="About 63% of positions, 751 in all, are non-moving."))
    assert invented_numbers(narrative, facts) == {"63", "751"}


def test_the_template_never_breaks_the_rules(facts) -> None:
    narrative = template_narrative(facts)
    assert invented_numbers(narrative, facts) == set()
    assert leaked_codes(narrative) == set()
    assert "Planned Overdue" in " ".join(narrative.findings), "the screens' word, not 'plan breach'"


def test_an_internal_code_in_the_prose_is_caught(facts) -> None:
    narrative = parse_reply(_reply(headline="Most exceptions are AWAITING_REQUESTER."))
    assert leaked_codes(narrative) == {"AWAITING_REQUESTER"}


def test_a_fenced_reply_parses_and_a_wrong_shape_does_not() -> None:
    assert parse_reply(f"```json\n{_reply()}\n```").headline.startswith("750")
    assert parse_reply("Here is your report: none") is None
    assert parse_reply(json.dumps({"headline": "x"})) is None, "findings are required"


# --- the narrative and its fallbacks ----------------------------------------------------


class _Provider:
    name = "fake"

    def __init__(self, *replies: str) -> None:
        self.replies = list(replies)
        self.prompts: list[str] = []

    def complete(self, messages, **kwargs) -> Completion:
        self.prompts.append(messages[0].content)
        return Completion(text=self.replies.pop(0), model=kwargs.get("model") or "gpt-4o", usage=TokenUsage())

    def check_connection(self) -> None:
        return None


class _Failing(_Provider):
    def complete(self, messages, **kwargs) -> Completion:
        raise AITimeoutError("no answer")


def _real_settings() -> Settings:
    return Settings(
        _env_file=None, llm_provider="openai", llm_base_url="https://example.invalid", llm_api_key="k", llm_model="m"
    )


def _use(monkeypatch: pytest.MonkeyPatch, provider) -> None:
    from app.core import prompts as prompt_module

    monkeypatch.setattr(prompt_module, "get_llm", lambda: provider)


def test_the_stub_is_skipped_for_the_template(facts) -> None:
    narrative = write_narrative(facts, settings=Settings(_env_file=None))
    assert narrative.source == DETERMINISTIC
    assert narrative.notes


def test_a_clean_reply_is_used_with_its_provenance(facts, monkeypatch) -> None:
    provider = _Provider(_reply())
    _use(monkeypatch, provider)
    narrative = write_narrative(facts, settings=_real_settings())

    assert narrative.source == MODEL
    assert narrative.prompt_id == PROMPT_ID and narrative.prompt_version == 1
    assert "{facts_json}" not in provider.prompts[0]
    assert '"non_moving_pct": 62.5' in provider.prompts[0]
    assert "NOT YOURS TO CHANGE" in provider.prompts[0]
    assert '"Requester not identified"' in provider.prompts[0], "the glossary reaches the model"


def test_an_invented_number_is_asked_again_then_accepted_when_fixed(facts, monkeypatch) -> None:
    provider = _Provider(_reply(headline="751 positions are idle."), _reply())
    _use(monkeypatch, provider)
    narrative = write_narrative(facts, settings=_real_settings())

    assert narrative.source == MODEL
    assert "751" in provider.prompts[1], "the second ask names the invented number"


def test_a_reply_with_a_code_is_asked_again(facts, monkeypatch) -> None:
    provider = _Provider(_reply(headline="Most are NON_MOVING."), _reply())
    _use(monkeypatch, provider)
    assert write_narrative(facts, settings=_real_settings()).source == MODEL
    assert "NON_MOVING" in provider.prompts[1]


def test_two_bad_replies_fall_back_to_the_template(facts, monkeypatch) -> None:
    _use(monkeypatch, _Provider("not json", _reply(headline="999 positions.")))
    assert write_narrative(facts, settings=_real_settings()).source == DETERMINISTIC


def test_a_provider_failure_falls_back_to_the_template(facts, monkeypatch) -> None:
    _use(monkeypatch, _Failing())
    narrative = write_narrative(facts, settings=_real_settings())
    assert narrative.source == DETERMINISTIC
    assert "did not answer" in narrative.notes[0]


# --- export ---------------------------------------------------------------------------------


def test_the_workbook_has_a_sheet_per_section(facts) -> None:
    report = {"facts": facts, "narrative": template_narrative(facts).to_dict()}
    book = load_workbook(build_report_workbook({"generated_at": "2026-10-13 10:00 UTC"}, report))
    assert book.sheetnames == ["Summary", "Key figures", "Received not issued", "Oldest exceptions"]
    assert book["Received not issued"].max_row == 3


# --- routes ------------------------------------------------------------------------------------


@pytest.fixture
def client(factory, db, monkeypatch: pytest.MonkeyPatch):
    def _client(enabled: bool) -> TestClient:
        monkeypatch.setenv("I13_REPORT_ENABLED", "true" if enabled else "false")
        monkeypatch.setenv("LLM_PROVIDER", "stub")
        get_settings.cache_clear()
        app = create_app()

        def _db():
            with factory() as session:
                yield session

        app.dependency_overrides[get_db] = _db
        app.dependency_overrides[require_snapshot] = lambda: SNAPSHOT
        # No `with`: entering the context runs the lifespan against the real database.
        return TestClient(app)

    yield _client
    get_settings.cache_clear()


def test_switched_off_every_route_is_404(client) -> None:
    c = client(enabled=False)
    assert c.post("/api/i13/reports/summary").status_code == 404
    assert c.get("/api/i13/reports").status_code == 404
    assert c.get("/api/i13/reports/abc").status_code == 404


def test_generate_list_open_and_export(client) -> None:
    c = client(enabled=True)
    created = c.post("/api/i13/reports/summary", headers={"X-Actor-Id": "chinmay"})
    assert created.status_code == 201
    report = created.json()
    assert report["generated_by"] == "chinmay"
    assert report["narrative_source"] == DETERMINISTIC
    assert report["facts"]["positions"]["total"] == 1200

    listed = c.get("/api/i13/reports").json()
    assert [r["id"] for r in listed] == [report["id"]]

    assert c.get(f"/api/i13/reports/{report['id']}").json()["headline"] == report["headline"]

    exported = c.get(f"/api/i13/reports/{report['id']}/export.xlsx")
    assert exported.status_code == 200
    assert "oar-utilization-report-all-plants-2026-10-13.xlsx" in exported.headers["content-disposition"]
    load_workbook(BytesIO(exported.content))


def test_an_unknown_report_is_404(client) -> None:
    assert client(enabled=True).get("/api/i13/reports/nope").status_code == 404
