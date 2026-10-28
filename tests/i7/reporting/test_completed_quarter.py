"""A completed quarter with three populated months -- the one thing the shared database cannot show yet.

Today the shared database holds October's recommendations only, so the monthly charts, the trend and the quarter-end
position cannot be seen for a *finished* quarter there. These tests build that quarter **synthetically, inside an
isolated in-memory database that exists only for the test** (never the shared one, and no historical recommendation is
ever invented anywhere else), run the real generator against it as of the day after quarter-end, and check what each
part of the report says. The month-end snapshots are captured by the real capture code with an injected date, exactly
as the nightly job would have on 2 Nov, 2 Dec and 2 Jan.
"""

from __future__ import annotations

from datetime import date
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.initiatives.i7.reporting import snapshots
from app.initiatives.i7.reporting.report_export import build_quarterly_report_workbook
from app.initiatives.i7.reporting.service import generate_quarterly_report
from app.schemas.i7.reports import AvailabilityStatus, PeriodStatus
from tests.i7.reporting.quarterly_fixtures import DEC, NOV, OCT, make_factory, seed_q4_2026, stub_stockout

JAN_1 = date(2027, 1, 1)
OCT_27 = date(2026, 10, 27)


@pytest.fixture
def session():
    factory = make_factory()
    with factory() as s:
        seed_q4_2026(s)
        yield s


def row(management, key):
    return next(r for r in management.monthly_activity.rows if r.key == key)


def capture_month_ends(session, monkeypatch) -> None:
    """What the nightly job would have saved: each month, captured in the first days of the next."""
    for month, captured_on, usage_built in (
        (OCT, date(2026, 11, 2), date(2026, 10, 30)),
        (NOV, date(2026, 12, 2), date(2026, 11, 30)),
        (DEC, date(2027, 1, 2), date(2026, 12, 30)),
    ):
        stub_stockout(monkeypatch, usage_built, at_risk=5, evaluated=50, insufficient=2)
        snapshots.capture_month(session, month, today=captured_on)


# --- the monthly charts, from three populated months ------------------------------------------


@pytest.fixture(autouse=True)
def _stockout(monkeypatch):
    stub_stockout(monkeypatch)


def test_every_month_of_the_completed_quarter_has_its_own_activity(session):
    report = generate_quarterly_report(session, "Q4 2026", today=JAN_1)
    management = report.management

    assert management.period.status is PeriodStatus.COMPLETE
    assert [(m.label, m.state, m.has_consumption_data, m.usable) for m in management.period.months] == [
        ("Oct 2026", "ended", True, True), ("Nov 2026", "ended", True, True), ("Dec 2026", "ended", True, True),
    ]

    generated = row(management, "recommendations_generated")
    assert generated.months == [3, 2, 2]  # Oct: A, B, C   Nov: A, D   Dec: A, E
    assert generated.quarter_total == 7 and generated.reconciles is True
    assert row(management, "submitted_for_approval").months == [1, 0, 0]
    assert row(management, "approval_decisions").months == [0, 0, 1]
    assert report.recommendations.total == 7
    assert management.status_distribution.total == 7 and management.status_distribution.measured_on == JAN_1


def test_the_same_data_read_in_the_middle_of_the_quarter_is_an_in_progress_report(session):
    management = generate_quarterly_report(session, "Q4 2026", today=OCT_27).management
    assert management.period.status is PeriodStatus.IN_PROGRESS
    assert management.period.cutoff == OCT_27
    assert [m.state for m in management.period.months] == ["in_progress", "not_started", "not_started"]
    assert "still in progress" in management.period.note


# --- the trend and the quarter-end position, from saved month-end snapshots -------------------


def test_without_snapshots_the_completed_quarter_says_so_and_falls_back_to_a_current_calculation(session):
    management = generate_quarterly_report(session, "Q4 2026", today=JAN_1).management
    assert [p.state for p in management.trends.months] == ["MISSING"] * 3
    assert management.trends.quarter_end.status is AvailabilityStatus.NOT_AVAILABLE
    assert management.position_basis == "current_calculation"
    assert management.data_coverage.basis == "current_calculation"


def test_with_three_month_end_snapshots_the_trend_has_three_points_and_the_position_is_the_december_one(session, monkeypatch):
    capture_month_ends(session, monkeypatch)
    stub_stockout(monkeypatch, date(2026, 12, 30), at_risk=5, evaluated=50, insufficient=2)

    management = generate_quarterly_report(session, "Q4 2026", today=JAN_1).management

    assert [p.label for p in management.trends.months] == ["Oct 2026", "Nov 2026", "Dec 2026"]
    assert all(p.state != "MISSING" for p in management.trends.months)
    assert all(p.month_end_valid for p in management.trends.months)  # each captured within days of its month's end
    end = management.trends.quarter_end
    assert (end.status, end.label) == (AvailabilityStatus.AVAILABLE, "Dec 2026")
    assert management.position_basis == "snapshot"  # the cards and per-plant charts read the one saved measurement
    assert management.data_coverage.basis == "snapshot"
    assert any("Dec 2026 month-end snapshot" in w for w in management.data_coverage.warnings)


def test_snapshots_are_final_and_do_not_move_when_the_report_is_generated_again(session, monkeypatch):
    capture_month_ends(session, monkeypatch)
    first = generate_quarterly_report(session, "Q4 2026", today=JAN_1).management.trends.model_dump_json()
    second = generate_quarterly_report(session, "Q4 2026", today=JAN_1).management.trends.model_dump_json()
    assert first == second


# --- the export -------------------------------------------------------------------------------


def test_the_excel_export_of_the_completed_quarter_carries_the_three_months(session, monkeypatch):
    capture_month_ends(session, monkeypatch)
    report = generate_quarterly_report(session, "Q4 2026", today=JAN_1)
    workbook = load_workbook(BytesIO(build_quarterly_report_workbook(report).getvalue()))
    text = " ".join(
        str(cell.value) for sheet in workbook for sheet_row in sheet.iter_rows() for cell in sheet_row if cell.value is not None
    )
    for label in ("Oct 2026", "Nov 2026", "Dec 2026"):
        assert label in text
