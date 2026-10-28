"""``python -m app.reporting.preview_quarterly``: the real report for a simulated date, exported, never saved.

Offline against an isolated in-memory database. The properties that matter here are the negative ones -- *nothing is
written* and *no part of generation reads the real clock* -- so most tests assert an absence.
"""

from __future__ import annotations

import ast
import datetime as dt
import inspect
import json
import sys
from datetime import date
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select

from app.initiatives.i7.reporting.period import latest_closed_quarter
from app.initiatives.i7.reporting.service import generate_quarterly_report
from app.models.i7_features import MaterialFeature
from app.models.i7_kpi_snapshot import KpiSnapshot, KpiSnapshotLog
from app.models.i7_recommendation import ApprovalLedgerEntry, Recommendation
from app.models.i7_reporting import QuarterlyReportRecord
from app.reporting import preview_quarterly
from app.reporting.preview_quarterly import ReadOnlyViolation, fence_read_only, main
from tests.i7.reporting.quarterly_fixtures import UTC, make_factory, rec, seed_q4_2026, stub_stockout

WATCHED = (Recommendation, ApprovalLedgerEntry, QuarterlyReportRecord, KpiSnapshot, KpiSnapshotLog, MaterialFeature)


@pytest.fixture(autouse=True)
def _stockout(monkeypatch):
    stub_stockout(monkeypatch)


@pytest.fixture
def factory(monkeypatch):
    f = make_factory()
    with f() as session:
        seed_q4_2026(session)
    monkeypatch.setattr(preview_quarterly, "get_sessionmaker", lambda: f)
    monkeypatch.setattr(preview_quarterly, "configure_logging", lambda settings: None)
    return f


def fingerprint(factory) -> dict:
    """Row counts of every table the report reads or could write, plus the recommendations' own content."""
    with factory() as session:
        counts = {m.__tablename__: session.execute(select(func.count()).select_from(m)).scalar_one() for m in WATCHED}
        counts["recommendation_rows"] = [
            (r.id, r.recommendation_id, r.status, r.generated_at, r.recommended_rop)
            for r in session.execute(select(Recommendation).order_by(Recommendation.id)).scalars()
        ]
    return counts


# --- what it does -----------------------------------------------------------------------------


def test_a_preview_as_of_new_year_exports_json_and_excel_and_saves_nothing(factory, tmp_path, capsys):
    before = fingerprint(factory)
    json_path, xlsx_path = tmp_path / "q4.json", tmp_path / "q4.xlsx"

    assert main(["--as-of", "2027-01-01", "--json", str(json_path), "--xlsx", str(xlsx_path)]) == 0

    out = capsys.readouterr().out
    assert out.startswith("PREVIEW -- nothing will be saved. Simulated date: 2027-01-01")
    assert "would target Q4 2026: DUE" in out
    assert out.rstrip().endswith("Done. Nothing was saved.")

    report = json.loads(json_path.read_text(encoding="utf-8"))
    assert report["metadata"]["quarter"] == "Q4 2026"
    assert report["metadata"]["generated_at"].startswith("2027-01-01")  # the simulated day, not the real one
    assert report["management"]["period"]["status"] == "COMPLETE"
    assert report["recommendations"]["total"] == 7

    workbook = load_workbook(BytesIO(xlsx_path.read_bytes()))
    assert "Metadata" in workbook.sheetnames and len(workbook.sheetnames) > 5

    assert fingerprint(factory) == before  # not one row added, changed or removed
    assert before["i7_quarterly_report"] == 0  # and in particular no report was saved


def test_the_quarter_defaults_to_the_latest_closed_one_as_of_the_simulated_day(factory, capsys):
    assert main(["--as-of", "2026-10-27"]) == 0
    out = capsys.readouterr().out
    assert "quarter ............ Q3 2026" in out and "OUTSIDE_WINDOW" in out  # today: Q3, and the nightly step would skip it
    assert main(["--as-of", "2027-01-01"]) == 0
    assert "quarter ............ Q4 2026" in capsys.readouterr().out


def test_a_running_quarter_previews_as_in_progress_not_complete(factory, tmp_path):
    path = tmp_path / "running.json"
    assert main(["--quarter", "Q4 2026", "--as-of", "2026-10-27", "--json", str(path)]) == 0
    period = json.loads(path.read_text(encoding="utf-8"))["management"]["period"]
    assert period["status"] == "IN_PROGRESS"
    assert period["cutoff"] == "2026-10-27"
    assert [m["state"] for m in period["months"]] == ["in_progress", "not_started", "not_started"]


def test_with_no_output_flags_it_still_runs_and_prints_the_summary_only(factory, capsys):
    assert main(["--as-of", "2027-01-01"]) == 0
    assert "period status ...... COMPLETE" in capsys.readouterr().out


def test_bad_arguments_are_usage_errors(factory, capsys):
    with pytest.raises(SystemExit) as caught:
        main(["--as-of", "2027-13-01"])
    assert caught.value.code == 2
    assert main(["--quarter", "Q5 2026", "--as-of", "2027-01-01"]) == 2
    assert "Invalid --quarter value" in capsys.readouterr().err


def test_a_generation_failure_exits_1_and_still_saves_nothing(factory, monkeypatch, capsys):
    before = fingerprint(factory)

    def boom(*args, **kwargs):
        raise RuntimeError("database blip")

    monkeypatch.setattr(preview_quarterly, "generate_quarterly_report", boom)
    assert main(["--as-of", "2027-01-01"]) == 1
    assert "Preview failed while generating: RuntimeError: database blip" in capsys.readouterr().err
    assert fingerprint(factory) == before


# --- it cannot write --------------------------------------------------------------------------


def test_the_module_never_imports_save_report_and_never_commits():
    tree = ast.parse(inspect.getsource(preview_quarterly))
    imported = {alias.name for node in ast.walk(tree) if isinstance(node, ast.ImportFrom) for alias in node.names}
    called = {
        node.func.attr if isinstance(node.func, ast.Attribute) else getattr(node.func, "id", "")
        for node in ast.walk(tree) if isinstance(node, ast.Call)
    }
    assert "save_report" not in imported | called
    assert not {"commit", "add", "add_all", "merge", "delete", "flush"} & called


def test_the_session_is_fenced_so_a_flush_or_a_commit_raises_and_reads_still_work(factory):
    with factory() as session:
        fence_read_only(session)
        assert session.execute(select(func.count()).select_from(Recommendation)).scalar_one() == 7  # reading is fine

        session.add(rec("Z", dt.datetime(2026, 11, 1, tzinfo=UTC)))
        with pytest.raises(ReadOnlyViolation):
            session.flush()
        session.rollback()
        with pytest.raises(ReadOnlyViolation):
            session.commit()
    assert fingerprint(factory)["i7_recommendation"] == 7


def test_even_a_generation_that_tried_to_write_would_fail_loudly_and_change_nothing(factory, monkeypatch, capsys):
    """The guarantee does not rest on today's generation happening to be read-only."""
    before = fingerprint(factory)

    def writes_a_row(session, quarter, **kwargs):
        session.add(QuarterlyReportRecord(
            quarter=quarter, period_start=date(2026, 10, 1), period_end=date(2026, 12, 31),
            generated_at=dt.datetime(2027, 1, 1, tzinfo=UTC), status="COMPLETED", report_version="1.2",
        ))
        session.flush()

    monkeypatch.setattr(preview_quarterly, "generate_quarterly_report", writes_a_row)
    assert main(["--as-of", "2027-01-01"]) == 1
    assert "read-only" in capsys.readouterr().err
    assert fingerprint(factory) == before


# --- the simulated date reaches everything ----------------------------------------------------


def _guard_the_wall_clock(monkeypatch) -> list[str]:
    """Make ``date.today()`` / ``datetime.now()`` raise in every module the report is built from."""

    class GuardDate(dt.date):
        @classmethod
        def today(cls):
            raise AssertionError("the real clock was read: date.today()")

    class GuardDatetime(dt.datetime):
        @classmethod
        def now(cls, tz=None):
            raise AssertionError("the real clock was read: datetime.now()")

        @classmethod
        def today(cls):
            raise AssertionError("the real clock was read: datetime.today()")

        @classmethod
        def utcnow(cls):
            raise AssertionError("the real clock was read: datetime.utcnow()")

    guarded = []
    for name, module in list(sys.modules.items()):
        if module is None or not name.startswith(("app.initiatives.i7.reporting", "app.schemas.i7.reports")):
            continue
        if getattr(module, "date", None) is dt.date:
            monkeypatch.setattr(module, "date", GuardDate)
            guarded.append(f"{name}.date")
        if getattr(module, "datetime", None) is dt.datetime:
            monkeypatch.setattr(module, "datetime", GuardDatetime)
            guarded.append(f"{name}.datetime")
    return guarded


def test_the_guard_itself_works(monkeypatch):
    """Without this, a green result below could just mean the guard guarded nothing."""
    guarded = _guard_the_wall_clock(monkeypatch)
    assert "app.initiatives.i7.reporting.period.date" in guarded
    assert "app.initiatives.i7.reporting.service.date" in guarded and "app.initiatives.i7.reporting.service.datetime" in guarded
    with pytest.raises(AssertionError, match="real clock"):
        latest_closed_quarter()  # no date given: falls back to date.today()


@pytest.mark.parametrize("as_of,quarter", [(date(2027, 1, 1), "Q4 2026"), (date(2026, 10, 27), "Q4 2026"), (date(2026, 10, 27), "Q3 2026")])
def test_no_part_of_generation_reads_the_real_clock_when_a_date_is_given(factory, monkeypatch, as_of, quarter):
    guarded = _guard_the_wall_clock(monkeypatch)
    assert len(guarded) >= 8  # service, period, quarter_management, monthly_activity, data_coverage, ... all watched

    now = dt.datetime(as_of.year, as_of.month, as_of.day, tzinfo=UTC)
    with factory() as session:
        report = generate_quarterly_report(session, quarter, today=as_of, generated_at=now)  # raises if anything peeks

    assert report.metadata.generated_at == now
    assert report.management.status_distribution.measured_on == as_of
    assert report.management.period.cutoff == min(date(2026, 12, 31) if quarter == "Q4 2026" else date(2026, 9, 30), as_of)
