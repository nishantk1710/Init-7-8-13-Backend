"""The quarterly report's Excel export: the existing sheets are kept, the management view is added, and every figure on it is
the saved report's own.

A real report is generated end to end on SQLite (every I07 table), then exported and read back with openpyxl -- so what is
checked is the workbook a person opens, against the report the page and the PDF are drawn from.
"""

from __future__ import annotations

from datetime import date, datetime, timezone
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import MetaData, create_engine
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

import app.models  # noqa: F401  (registers every model on Base.metadata)
from app.initiatives.i7.reporting import kpi_definitions, snapshots
from app.initiatives.i7.reporting import service as report_service
from app.initiatives.i7.reporting.report_export import NOT_AVAILABLE, build_quarterly_report_workbook
from app.initiatives.i7.reporting.service import generate_quarterly_report
from app.initiatives.i7.stockout_risk.counts import StockoutRiskCounts
from app.models.base import Base
from app.models.i7_recommendation import ApprovalLedgerEntry, Recommendation
from app.models.i7_staging import StagedConsumption

ORIGINAL_SHEETS = ["Metadata", "Baseline Comparison", "Scope & Data Quality", "Limitations"]
TECHNICAL_SHEETS = ["Demand Classification", "Forecast Accuracy", "OAR", "Recommendations", "Approval Activity"]
MANAGEMENT_SHEETS = [
    "Period", "KPI Cards", "Monthly Activity", "Position", "Status Distribution", "Forecast vs Actual", "SAP Adoption",
    "Trends", "Data Coverage", "Definitions",
]


@pytest.fixture
def db():
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    metadata = MetaData()
    for name, table in Base.metadata.tables.items():
        if name.startswith("i7_"):
            table.to_metadata(metadata).indexes.clear()
    metadata.create_all(engine)
    with Session(engine) as session:
        yield session


AFTER_Q4 = date(2027, 1, 15)  # Q4 2026 is complete


@pytest.fixture(autouse=True)
def fixed_today(monkeypatch):
    """The report is judged as of a fixed day, so the quarter is complete and the tests do not depend on the clock."""

    class FixedDate(date):
        @classmethod
        def today(cls):
            return AFTER_Q4

    monkeypatch.setattr(report_service, "date", FixedDate)


@pytest.fixture(autouse=True)
def stockout_never_built(monkeypatch):
    monkeypatch.setattr(
        kpi_definitions, "stockout_risk_counts",
        lambda session, base, **kw: StockoutRiskCounts(at_risk=0, evaluated=0, insufficient_data=0, as_of=None),
    )


def at(year, month, day=10):
    return datetime(year, month, day, 9, 0, tzinfo=timezone.utc)


def seed(db):
    db.add_all([
        Recommendation(
            recommendation_id=f"REC-{m}-1300-1", sap_material_number=m, sap_plant_code="1300", feature_run_id=1, forecast_run_id=1,
            inventory_run_id=1, oar_run_id=1, policy_id="p", policy_version=1, formula_version="f", status=status, impact_status="X",
            generated_at=at(2026, month), unit_price=Decimal("10"), current_rop=2, recommended_rop=12, current_max_stock=100,
            recommended_max_stock=60, current_safety_stock=1, recommended_safety_stock=3,
        )
        for m, month, status in (("A", 10, "READY_FOR_REVIEW"), ("B", 10, "PENDING_APPROVAL"), ("C", 11, "APPROVED"))
    ])
    db.add_all([
        ApprovalLedgerEntry(recommendation_id="REC-C-1300-1", recommendation_version=1, actor_id="u", actor_role="End User", action=action,
                            previous_status="READY_FOR_REVIEW", new_status="PENDING_APPROVAL", timestamp=when)
        for action, when in (("SUBMIT", at(2026, 11, 12)), ("APPROVE", at(2026, 11, 20)))
    ])
    db.add(StagedConsumption(sap_material_number="A", sap_plant_code="1300", period=date(2026, 10, 1), quantity=Decimal("5"),
                             unit_of_measure="EA", movement_count=3, issue_count=3, reversal_count=0, staging_run_id=1, source_table="n_mseg"))
    db.commit()


def generated(db, *, capture=False):
    seed(db)
    if capture:
        snapshots.capture_month(db, date(2026, 10, 1), today=date(2026, 11, 3))
    return generate_quarterly_report(db, "Q4 2026")


def workbook(report):
    return load_workbook(BytesIO(build_quarterly_report_workbook(report).getvalue()))


def rows(sheet):
    return [list(r) for r in sheet.iter_rows(values_only=True)]


def facts_of(sheet, start=1, count=None):
    """The (field, value) pairs of a sheet's leading facts table: the first two cells of each row."""
    pairs = [tuple(r[:2]) for r in rows(sheet)[start:]]
    out = {}
    for key, value in pairs:
        if key is None:
            break
        out[key] = value
        if count is not None and len(out) == count:
            break
    return out


def table_after(sheet, header_first_cell):
    """The rows of the table whose header row starts with ``header_first_cell``, up to the next blank row."""
    all_rows = rows(sheet)
    start = next(i for i, r in enumerate(all_rows) if r and r[0] == header_first_cell)
    out = []
    for r in all_rows[start + 1:]:
        if all(v is None for v in r):
            break
        out.append(r)
    return all_rows[start], out


# --- the sheets ------------------------------------------------------------------------------


def test_the_existing_sheets_come_first_and_unchanged_in_name_and_order(db):
    names = workbook(generated(db)).sheetnames
    assert names[:4] == ORIGINAL_SHEETS


def test_a_current_report_has_the_technical_and_management_sheets(db):
    names = workbook(generated(db)).sheetnames
    for name in TECHNICAL_SHEETS + MANAGEMENT_SHEETS:
        assert name in names, name


def test_a_report_saved_before_the_management_view_exports_without_those_sheets(db):
    old = generated(db).model_copy(update={"management": None})
    names = workbook(old).sheetnames
    assert names[:4] == ORIGINAL_SHEETS
    assert not set(MANAGEMENT_SHEETS) & set(names)
    assert set(TECHNICAL_SHEETS) <= set(names)


def test_the_existing_sheets_still_hold_what_they_held(db):
    report = generated(db)
    wb = workbook(report)
    meta = facts_of(wb["Metadata"])
    assert meta["Quarter"] == "Q4 2026" and meta["Report Version"] == "1.2"
    baseline = table_after(wb["Baseline Comparison"], "Metric")[1]
    assert [r[0] for r in baseline] == ["Safety Stock", "ROP", "Max Stock", "Lead Time"]
    assert [r[0] for r in rows(wb["Limitations"])[1:]] == report.limitations.items


# --- the figures are the report's own ----------------------------------------------------------


def test_every_kpi_card_on_the_sheet_is_the_reports_card(db):
    report = generated(db)
    header, body = table_after(workbook(report)["KPI Cards"], "Card")
    col = {name: i for i, name in enumerate(header)}
    assert [r[col["Card"]] for r in body] == [c.title for c in report.management.kpis]
    for row, card in zip(body, report.management.kpis):
        assert row[col["Unit"]] == card.unit
        assert row[col["Kind"]] == card.kind.value
        assert row[col["Basis"]] == card.basis
        assert row[col["Covers"]] == card.period_label
        assert row[col["Denominator"]] == (card.denominator if card.denominator is not None else NOT_AVAILABLE)
        if card.status.value == "AVAILABLE":
            assert Decimal(str(row[col["Value"]])) == card.value
        else:
            assert row[col["Value"]] == NOT_AVAILABLE  # never a blank or a 0


def test_a_kpi_that_has_no_value_says_why_in_the_export(db):
    report = generated(db)
    header, body = table_after(workbook(report)["KPI Cards"], "Card")
    col = {name: i for i, name in enumerate(header)}
    stockout = next(r for r in body if r[col["Card"]] == "Critical stockout risk")
    assert stockout[col["Value"]] == NOT_AVAILABLE
    assert stockout[col["Status"]] == "NOT_AVAILABLE"
    assert stockout[col["Measured on"]] == NOT_AVAILABLE or stockout[col["Why not available"]] != ""


def test_monthly_activity_has_the_three_months_the_totals_and_whether_they_add_up(db):
    report = generated(db)
    header, body = table_after(workbook(report)["Monthly Activity"], "Metric")
    assert header[3:6] == ["Oct 2026", "Nov 2026", "Dec 2026"]
    by_label = {r[0]: r for r in body}
    recs = by_label["Recommendations generated"]
    assert recs[3:7] == [2, 1, 0, 3]  # Oct, Nov, Dec, Quarter total
    assert recs[header.index("Months add up to the total")] == "Yes"
    # A month that has not started or is not covered is the words, not a blank or a zero.
    consumption = by_label["SAP consumption"]
    assert consumption[3] == 3 and consumption[4] == NOT_AVAILABLE
    assert by_label["Confirmed SAP adoptions"][3:7] == [NOT_AVAILABLE] * 4


def test_decisions_come_from_dated_events(db):
    header, body = table_after(workbook(generated(db))["Monthly Activity"], "Metric")
    by_label = {r[0]: r for r in body}
    assert by_label["Sent for approval"][3:7] == [0, 1, 0, 1]
    assert by_label["Approval decisions"][3:7] == [0, 1, 0, 1]


def test_the_position_sheet_has_every_plant_and_says_when_the_stockout_was_not_measured(db):
    report = generated(db)
    wb = workbook(report)
    header, body = table_after(wb["Position"], "Plant")
    assert [r[0] for r in body] == ["All plants" if p.plant is None else p.plant for p in report.management.plant_positions]
    all_plants = body[0]
    assert all_plants[header.index("Material-plants")] == report.management.position.population
    assert all_plants[header.index("Stockout measured as of")] == NOT_AVAILABLE  # the usage was never built here


def test_the_status_distribution_adds_up_to_its_total(db):
    wb = workbook(generated(db))
    facts = facts_of(wb["Status Distribution"])
    header, body = table_after(wb["Status Distribution"], "Status")
    assert sum(r[1] for r in body) == facts["Records"] == 3


def test_adoption_is_not_measured_and_the_export_says_so(db):
    facts = facts_of(workbook(generated(db))["SAP Adoption"])
    assert facts["Status"] == "UNKNOWN"
    assert facts["Confirmed adoptions"] == NOT_AVAILABLE
    assert "not an adoption" in facts["Reason"]


def test_data_coverage_lists_the_dates_and_every_warning(db):
    report = generated(db)
    wb = workbook(report)
    facts = facts_of(wb["Data Coverage"])
    assert facts["Recommendations counted"] == 3
    assert facts["Figures are"] == "Current calculations"
    warnings = [r[0] for r in table_after(wb["Data Coverage"], "Limitation or warning")[1]]
    assert warnings == report.management.data_coverage.warnings
    assert any("never as zero" in w for w in warnings)


# --- snapshots in the export -------------------------------------------------------------------


def test_the_trends_sheet_shows_a_captured_month_and_a_missing_one_and_the_unavailable_quarter_end(db):
    report = generated(db, capture=True)
    wb = workbook(report)
    header, body = table_after(wb["Trends"], "Month")
    col = {name: i for i, name in enumerate(header)}
    october, november, december = body
    assert october[col["Snapshot"]] == "FINAL" and october[col["Measured on"]] == "2026-11-03"
    assert october[col["Stockout reading usable"]] == "No"  # the usage was never built
    assert october[col["At stockout risk"]] == NOT_AVAILABLE
    assert october[col["Excess candidates"]] == 2  # A and B were in force on 31 Oct; C was generated in November
    assert november[col["Snapshot"]] == "MISSING" and november[col["Excess candidates"]] == NOT_AVAILABLE
    all_rows = rows(wb["Trends"])
    at_index = next(i for i, r in enumerate(all_rows) if r[0] == "Quarter-end position")
    quarter_end = facts_of(wb["Trends"], start=at_index + 2)
    assert quarter_end["Status"] == "NOT_AVAILABLE"
    assert "Dec 2026" in quarter_end["Why not available"]


def test_the_export_says_counts_are_never_added_across_months(db):
    texts = [v for r in rows(workbook(generated(db, capture=True))["Trends"]) for v in r if isinstance(v, str)]
    assert any("never added into a quarter total" in t for t in texts)


# --- the whole workbook ------------------------------------------------------------------------


def test_no_cell_in_the_management_sheets_is_an_unexplained_none_in_a_value_column(db):
    """Nothing is silently blank where a figure belongs: it is a number or the words 'Not available'."""
    wb = workbook(generated(db, capture=True))
    for name in ("KPI Cards", "Position", "Trends"):
        header, body = table_after(wb[name], "Card" if name == "KPI Cards" else "Plant" if name == "Position" else "Month")
        for row in body:
            for title, value in zip(header, row):
                if title in ("Value", "At stockout risk", "Excess opportunity", "Working capital impact (estimate)", "Material-plants"):
                    assert value is not None, f"{name} / {title}"


def test_the_export_is_a_valid_workbook_that_opens_and_has_no_empty_sheets(db):
    wb = workbook(generated(db, capture=True))
    for sheet in wb.worksheets:
        assert sheet.max_row >= 2, sheet.title
