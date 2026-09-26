"""Seed loader tests.

Mechanics are tested against tiny workbooks held in an in-memory Data Lake --
the real extracts are 800 MB and never go near a test run. One end-to-end test
loads a small real table when both the database and the extracts are available,
which is what catches "works on a synthetic 3-row sheet, falls over on SAP's
actual output".

The workbooks live in the fake storage rather than a temp directory because the
local filesystem adapter was removed: the system reads from
``stvziaicomnonprod`` and nothing else, and the tests exercise that path.
"""

from __future__ import annotations

import io

import openpyxl
import pytest
from sqlalchemy import text

from app.core.config import get_settings
from app.core.db import MSSQL, backend_of, get_engine, get_sessionmaker
from app.core.storage import Storage
from app.integrations.storage.adls import AzureDataLakeStorage
from app.models.ingestion import IngestionRun
from app.seed.loader import STATUS_SUCCEEDED, load_table
from app.seed.manifest import ALL_FILES, BY_TABLE, EXTRACTS, REPORTS, TABLES, spec_for
from app.seed.reader import (
    ExtractFormatError,
    column_name,
    read_headers,
    read_rows,
    to_text,
    unique_column_names,
)
from tests.fake_adls import FakeDataLakeServiceClient

FAKE_URL = "abfss://raw@stvziaicomnonprod.dfs.core.windows.net/extracts"


def _workbook_bytes(rows: list[list]) -> bytes:
    """A minimal .xlsx whose first row is the header."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def _storage_holding(key: str, rows: list[list]) -> Storage:
    """A Data Lake adapter containing one workbook at ``key``."""
    return _storage_holding_bytes(key, _workbook_bytes(rows))


def _storage_holding_bytes(key: str, payload: bytes) -> Storage:
    storage = AzureDataLakeStorage(FAKE_URL, service_client=FakeDataLakeServiceClient())
    with storage.open_write(key) as handle:
        handle.write(payload)
    return storage


def _azure_sql_configured() -> bool:
    url = get_settings().database_url
    return bool(url) and backend_of(url) == MSSQL


def _extracts_configured() -> bool:
    """Whether STORAGE_URL names a Data Lake we could actually read.

    Checks the scheme, not just emptiness. A leftover local path is set but
    unusable, and the difference between "skipped, not configured" and "failed"
    should not depend on a stale value in someone's .env.
    """
    return get_settings().storage_url.lower().startswith(("abfss://", "abfs://"))


needs_db = pytest.mark.skipif(
    not _azure_sql_configured(),
    reason="DATABASE_URL does not name an Azure SQL database (reachable only inside the VNet)",
)
needs_extracts = pytest.mark.skipif(
    not _extracts_configured(), reason="STORAGE_URL does not name an Azure Data Lake"
)


# --- Column naming --------------------------------------------------------


class TestColumnNaming:
    @pytest.mark.parametrize(
        ("header", "expected"),
        [
            ("Material", "material"),
            ("Ext. Material Group", "ext_material_group"),
            ("DF at client level", "df_at_client_level"),
            ("MRP Type", "mrp_type"),
            ("GR/IR clearing value in FC", "gr_ir_clearing_value_in_fc"),
            ("Cont. Non-Textile Parts", "cont_non_textile_parts"),
        ],
    )
    def test_business_labels_become_identifiers(self, header: str, expected: str) -> None:
        assert column_name(header, 1) == expected

    def test_blank_header_keeps_the_column(self) -> None:
        """An unnamed column still holds data; dropping it would lose rows' values."""
        assert column_name(None, 7) == "column_7"
        assert column_name("   ", 7) == "column_7"

    def test_leading_digit_is_prefixed(self) -> None:
        """'2nd License Qty' is a real EKPO header and is not a legal identifier."""
        assert column_name("2nd License Qty", 3).startswith("c_")

    def test_duplicate_headers_are_suffixed_not_collapsed(self) -> None:
        """16 of the 28 extracts have duplicate headers -- both columns must survive."""
        names = unique_column_names(["Amount", "Amount", "Amount"])
        assert names == ["amount", "amount_1", "amount_2"]
        assert len(set(names)) == 3

    def test_long_headers_are_truncated_to_the_postgres_limit(self) -> None:
        assert len(column_name("x" * 200, 1)) <= 63


# --- Value rendering ------------------------------------------------------


class TestToText:
    def test_none_and_blank_stay_none(self) -> None:
        assert to_text(None) is None
        assert to_text("   ") is None

    def test_integral_floats_lose_the_decimal(self) -> None:
        """Excel stores every number as a float; 2000000131.0 is a material number."""
        assert to_text(2000000131.0) == "2000000131"

    def test_non_integral_floats_are_preserved(self) -> None:
        assert to_text(12.5) == "12.5"

    def test_material_numbers_are_never_coerced(self) -> None:
        """The whole reason the raw layer is text."""
        assert to_text("000000008000000000") == "000000008000000000"

    def test_dates_become_iso(self) -> None:
        from datetime import datetime

        assert to_text(datetime(2025, 7, 31)) == "2025-07-31"
        assert to_text(datetime(2025, 7, 31, 14, 30)) == "2025-07-31 14:30:00"

    def test_booleans_use_the_sap_flag_convention(self) -> None:
        assert to_text(True) == "X"
        assert to_text(False) is None


# --- Reading --------------------------------------------------------------


class TestReader:
    def test_reads_headers_and_rows(self) -> None:
        storage = _storage_holding("T.XLSX", [["Material", "Plant"], ["1000", "1300"]])
        assert read_headers(storage, "T.XLSX") == ["material", "plant"]
        assert list(read_rows(storage, "T.XLSX")) == [("1000", "1300")]

    def test_fully_blank_rows_are_skipped(self) -> None:
        """Excel exports carry trailing empty rows inside the declared dimension."""
        storage = _storage_holding(
            "T.XLSX",
            [["Material", "Plant"], ["1000", "1300"], [None, None], ["2000", "1500"]],
        )
        assert list(read_rows(storage, "T.XLSX")) == [("1000", "1300"), ("2000", "1500")]

    def test_short_rows_are_padded_not_shifted(self) -> None:
        """A ragged row must not slide every value one column to the left."""
        storage = _storage_holding("T.XLSX", [["a", "b", "c"], ["1", "2", "3"], ["4"]])
        assert list(read_rows(storage, "T.XLSX")) == [("1", "2", "3"), ("4", None, None)]

    def test_header_mismatch_raises(self) -> None:
        """The guard that makes concatenating split files safe."""
        storage = _storage_holding("T.XLSX", [["Material", "Plant"], ["1000", "1300"]])
        with pytest.raises(ExtractFormatError, match="header does not match"):
            list(read_rows(storage, "T.XLSX", expected=["material", "werks"]))

    def test_reads_a_named_sheet_not_the_first(self) -> None:
        """BMM's ZMM065 data is the third sheet of five."""
        workbook = openpyxl.Workbook()
        workbook.active.title = "Pivot"
        workbook.active.append(["Row Labels", "Sum of Value"])
        data = workbook.create_sheet("Sheet1")
        data.append(["Mat.Code", "Plant"])
        data.append(["8000004187", "1300"])
        buffer = io.BytesIO()
        workbook.save(buffer)
        workbook.close()

        storage = _storage_holding_bytes("R.xlsx", buffer.getvalue())
        assert read_headers(storage, "R.xlsx", sheet="Sheet1") == ["mat_code", "plant"]
        assert list(read_rows(storage, "R.xlsx", sheet="Sheet1")) == [("8000004187", "1300")]

    def test_header_row_offset_skips_a_title_row(self) -> None:
        """Both ZMM065 reports have a blank row above the header."""
        storage = _storage_holding(
            "R.xlsx", [[None, None], ["Mat.Code", "Plant"], ["8000004187", "1300"]]
        )
        assert read_headers(storage, "R.xlsx", header_row=2) == ["mat_code", "plant"]
        assert list(read_rows(storage, "R.xlsx", header_row=2)) == [("8000004187", "1300")]

    def test_default_header_row_on_a_titled_sheet_is_useless(self) -> None:
        """Why header_row exists: row 1 yields column_1, column_2 and unusable labels."""
        storage = _storage_holding(
            "R.xlsx", [[None, None], ["Mat.Code", "Plant"], ["8000004187", "1300"]]
        )
        assert read_headers(storage, "R.xlsx") == ["column_1", "column_2"]

    def test_unknown_sheet_name_lists_what_is_there(self) -> None:
        storage = _storage_holding("R.xlsx", [["a"], ["1"]])
        with pytest.raises(ExtractFormatError, match="Sheets present"):
            read_headers(storage, "R.xlsx", sheet="Nope")

    def test_empty_workbook_raises(self) -> None:
        storage = _storage_holding("T.XLSX", [])
        with pytest.raises(ExtractFormatError):
            read_headers(storage, "T.XLSX")

    def test_a_workbook_is_read_through_the_data_lake_adapter(self) -> None:
        """The reader must work on a spooled download, not only a real file.

        openpyxl seeks inside the zip, so this is the property that would have
        broken had open_read handed back Azure's forward-only downloader.
        """
        storage = _storage_holding("T.XLSX", [["Material"], ["1000"]])
        assert isinstance(storage, AzureDataLakeStorage)
        assert list(read_rows(storage, "T.XLSX")) == [("1000",)]


# --- Manifest -------------------------------------------------------------


class TestManifest:
    def test_table_names_are_unique(self) -> None:
        names = [spec.table for spec in EXTRACTS]
        assert len(names) == len(set(names))

    def test_no_file_is_loaded_into_two_tables(self) -> None:
        files = [key for spec in EXTRACTS for key in spec.files]
        assert len(files) == len(set(files))

    def test_raw_prefix_is_applied(self) -> None:
        assert spec_for("marc").raw_table == "raw_marc"

    def test_unknown_table_names_the_alternatives(self) -> None:
        with pytest.raises(KeyError, match="Known tables"):
            spec_for("nope")

    def test_split_tables_declare_their_files_in_order(self) -> None:
        assert BY_TABLE["mseg"].files == (f"{TABLES}/Mseg_1.XLSX", f"{TABLES}/MSEG_2.XLSX")
        assert BY_TABLE["cdhdr"].files == (f"{TABLES}/CDHDR1.XLSX", f"{TABLES}/CDHDR2.XLSX")

    def test_keys_are_prefixed_by_delivery_folder(self) -> None:
        """One storage root spans both deliveries, so every key carries its prefix."""
        for spec in EXTRACTS:
            for key in spec.files:
                assert key.startswith((TABLES, REPORTS)), key
                assert "\\" not in key, f"{key}: storage keys use forward slashes"

    def test_reports_declare_their_sheet_and_header_row(self) -> None:
        """The reports are multi-sheet with title rows; defaults would read garbage."""
        assert BY_TABLE["zmm065_bmm"].sheet == "Sheet1"
        assert BY_TABLE["zmm065_bmm"].header_row == 2
        assert BY_TABLE["zmm065_gb"].sheet == "Sheet2"
        assert BY_TABLE["zmm065_gb"].header_row == 2
        assert BY_TABLE["gr_30day"].sheet == "GR REPORT"

    def test_table_extracts_use_the_defaults(self) -> None:
        """Only the reports need overrides -- the SAP dumps are sheet 1, row 1."""
        for name in ("mara", "marc", "mseg", "cdhdr"):
            assert BY_TABLE[name].sheet is None
            assert BY_TABLE[name].header_row == 1

    @needs_extracts
    def test_every_manifest_file_exists_in_storage(self) -> None:
        """Catches a renamed or missing extract before a load run does."""
        from app.core.storage import get_storage

        available = set(get_storage().list())
        assert ALL_FILES <= available, f"missing: {sorted(ALL_FILES - available)}"


# --- Loading --------------------------------------------------------------


@needs_db
class TestLoading:
    def test_loads_a_table_and_records_the_run(self) -> None:
        from app.seed.manifest import ExtractSpec

        storage = _storage_holding(
            "Thing.XLSX",
            [["Material", "Plant"], ["1000000000", "1300"], ["2000000131", "1500"]],
        )
        spec = ExtractSpec(table="seedtest", files=("Thing.XLSX",), sap_table="ZTEST")

        import app.seed.loader as loader

        original = loader.get_storage
        loader.get_storage = lambda: storage
        try:
            result = load_table(spec, force=True)
            assert result.status == STATUS_SUCCEEDED
            assert result.rows == 2

            with get_engine().connect() as connection:
                rows = connection.execute(
                    text("SELECT material, plant FROM raw_seedtest ORDER BY material")
                ).all()
            assert rows == [("1000000000", "1300"), ("2000000131", "1500")]

            with get_sessionmaker()() as session:
                runs = (
                    session.query(IngestionRun)
                    .filter(IngestionRun.target_table == "raw_seedtest")
                    .all()
                )
                assert len(runs) == 1
                assert runs[0].row_count == 2
                assert runs[0].source_sha256 and len(runs[0].source_sha256) == 64
        finally:
            loader.get_storage = original
            with get_engine().begin() as connection:
                connection.execute(text("DROP TABLE IF EXISTS raw_seedtest"))
            with get_sessionmaker()() as session:
                session.query(IngestionRun).filter(
                    IngestionRun.target_table == "raw_seedtest"
                ).delete()
                session.commit()

    def test_unchanged_source_is_skipped_then_forced(self) -> None:
        from app.seed.manifest import ExtractSpec

        storage = _storage_holding("Thing.XLSX", [["Material"], ["1000000000"]])
        spec = ExtractSpec(table="seedskip", files=("Thing.XLSX",), sap_table="ZTEST")

        import app.seed.loader as loader

        original = loader.get_storage
        loader.get_storage = lambda: storage
        try:
            assert load_table(spec, force=True).status == STATUS_SUCCEEDED
            assert load_table(spec).status == "skipped"
            assert load_table(spec, force=True).status == STATUS_SUCCEEDED
        finally:
            loader.get_storage = original
            with get_engine().begin() as connection:
                connection.execute(text("DROP TABLE IF EXISTS raw_seedskip"))
            with get_sessionmaker()() as session:
                session.query(IngestionRun).filter(
                    IngestionRun.target_table == "raw_seedskip"
                ).delete()
                session.commit()

    def test_a_failed_load_leaves_the_previous_table_intact(self) -> None:
        """The transaction boundary: a bad second file must not destroy good data."""
        from app.seed.manifest import ExtractSpec

        import app.seed.loader as loader

        storage = _storage_holding("Good.XLSX", [["Material"], ["1000000000"]])
        good = ExtractSpec(table="seedatomic", files=("Good.XLSX",), sap_table="ZTEST")

        original = loader.get_storage
        loader.get_storage = lambda: storage
        try:
            assert load_table(good, force=True).status == STATUS_SUCCEEDED

            # A second file whose header disagrees -- the load must abort.
            with storage.open_write("Bad.XLSX") as handle:
                handle.write(_workbook_bytes([["Different"], ["x"]]))
            broken = ExtractSpec(
                table="seedatomic", files=("Good.XLSX", "Bad.XLSX"), sap_table="ZTEST"
            )
            assert load_table(broken, force=True).status == "failed"

            with get_engine().connect() as connection:
                rows = connection.execute(text("SELECT material FROM raw_seedatomic")).all()
            assert rows == [("1000000000",)], "previous contents should survive a failed reload"
        finally:
            loader.get_storage = original
            with get_engine().begin() as connection:
                connection.execute(text("DROP TABLE IF EXISTS raw_seedatomic"))
            with get_sessionmaker()() as session:
                session.query(IngestionRun).filter(
                    IngestionRun.target_table == "raw_seedatomic"
                ).delete()
                session.commit()


@needs_db
@needs_extracts
def test_end_to_end_on_a_real_extract() -> None:
    """Load the smallest real workbook. Synthetic sheets do not catch SAP's quirks."""
    result = load_table(spec_for("eina"), force=True)
    assert result.status == STATUS_SUCCEEDED
    assert result.rows > 4000

    with get_engine().connect() as connection:
        material = connection.execute(
            text("SELECT TOP 1 material FROM raw_eina WHERE material IS NOT NULL")
        ).scalar_one()

    # The point of the text layer: SAP keys keep their exact form.
    assert isinstance(material, str)
    assert material.isdigit()
